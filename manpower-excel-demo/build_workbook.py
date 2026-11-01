"""Builds Manpower_Management_Demo.xlsx - a formula-driven Excel system for a
UAE manpower supply company (own workers + supplier workers deployed to clients).

Run:  python build_workbook.py
"""
import random
from datetime import date

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

random.seed(7)
OUT = "Manpower_Management_Demo.xlsx"

# ---------------------------------------------------------------- styles
FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True, color="1F3864")
F_SUB = Font(name=FONT, size=9, italic=True, color="595959")
F_BOLD = Font(name=FONT, size=10, bold=True)
FILL_HEAD = PatternFill("solid", fgColor="1F3864")
FILL_CALC_HEAD = PatternFill("solid", fgColor="2E75B6")
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")
FILL_TOTAL = PatternFill("solid", fgColor="D9E1F2")
FILL_RED = PatternFill("solid", fgColor="F8CBAD")
FILL_AMBER = PatternFill("solid", fgColor="FFE699")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
AED = '#,##0.00;(#,##0.00);"-"'
NUM = '#,##0;(#,##0);"-"'
DATE = "dd-mmm-yyyy"

HEADER_ROW = 4
FIRST = 5

# Settings cell references
S_COMPANY = "Settings!$B$4"
S_TRN = "Settings!$B$5"
S_MOHRE = "Settings!$B$6"
S_BANK = "Settings!$B$7"
S_START = "Settings!$B$8"
S_END = "Settings!$B$9"
S_DIV = "Settings!$B$10"
S_STD = "Settings!$B$11"
S_HRS = "Settings!$B$12"
S_OT = "Settings!$B$13"
S_VAT = "Settings!$B$14"
S_ALERT = "Settings!$B$15"

TRADES = ["Helper", "Mason", "Carpenter", "Electrician", "Plumber",
          "Steel Fixer", "Driver", "Cleaner", "Security Guard"]
DED_TYPES = ["Advance", "Fine", "Accommodation", "Other"]

wb = Workbook()


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    ws["A2"] = sub
    ws["A2"].font = F_SUB


def headers(ws, cols, row=HEADER_ROW, calc_from=None):
    """cols: list of (header, width). calc_from: 1-based col index where
    formula (calculated) columns start - shown in a lighter header colour."""
    for i, (h, w) in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = F_HEAD
        c.fill = FILL_CALC_HEAD if calc_from and i >= calc_from else FILL_HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 32
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def put(ws, r, c, v, fmt=None, inp=False, bold=False):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = F_INPUT if inp else (F_BOLD if bold else F_BASE)
    if fmt:
        cell.number_format = fmt
    cell.border = BORDER
    return cell


def add_list_validation(ws, rng, formula):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.error = "Please pick a value from the list"
    ws.add_data_validation(dv)
    dv.add(rng)


# ================================================================ README
ws = wb.active
ws.title = "README"
title(ws, "Manpower Management System - Excel Demo",
      "Own workers + supplier workers -> clients -> timesheet -> payroll, supplier payment, client invoice")
lines = [
    ("HOW THE FLOW WORKS", None),
    ("1. Settings", "Choose the payroll month and company rules (OT rate, VAT, working days)."),
    ("2. Clients / Suppliers", "Master list of customers and manpower suppliers (one row each)."),
    ("3. ClientRates", "What you CHARGE each client per hour, per trade (normal + OT)."),
    ("4. SupplierRates", "What each supplier CHARGES YOU per hour, per trade (normal + OT)."),
    ("5. Workers", "Every worker: Source = Own or Supplier. Own workers have salary, IBAN, labour card."),
    ("6. Timesheet", "Monthly hours received from the client. ONLY enter columns A-H; the rest is calculated."),
    ("7. Deductions", "Salary advances, fines, accommodation etc. for own workers."),
    ("", ""),
    ("AUTOMATIC OUTPUT SHEETS", None),
    ("Dashboard", "Month KPIs: billing, supplier cost, own payroll, profit, profit per client, document alerts."),
    ("Payroll", "Own workers' salary for the month (fixed, hourly, OT, deductions, net pay)."),
    ("Payslip", "Pick a Worker ID in cell C4 -> printable payslip."),
    ("ClientInvoice", "Pick a Client ID in cell C4 -> tax invoice with trade-wise hours and 5% VAT."),
    ("SupplierStatement", "How much to pay each supplier this month, incl. VAT and due date."),
    ("WPS_SIF", "Salary rows in WPS SIF layout (EDR + SCR). Confirm exact format with your bank/exchange."),
    ("", ""),
    ("COLOUR LEGEND", None),
    ("Blue text / yellow cell", "Input - type here."),
    ("Black text", "Formula - do not overwrite."),
    ("Dark blue header", "Input column.   Light blue header = calculated column."),
    ("Red row", "Visa / Emirates ID expired or expiring within the alert days (Settings)."),
    ("", ""),
    ("MONTHLY ROUTINE", None),
    ("Step 1", "Settings!B8 -> set the month (1st day of month)."),
    ("Step 2", "Paste/enter client timesheets in Timesheet (Month, Client, Site, Worker ID, Days, Hours, OT)."),
    ("Step 3", "Enter advances / fines in Deductions."),
    ("Step 4", "Check Dashboard -> print Payslips, ClientInvoice, SupplierStatement, export WPS_SIF."),
    ("", ""),
    ("ASSUMPTIONS (change in Settings)", None),
    ("Monthly salary", "Absent days = Standard working days - Days worked; deduction = (Basic+Allowances)/30 x absent days."),
    ("Overtime", "UAE practice: OT hourly = Basic/30/8 x 1.25 (multiplier editable). Hourly workers: rate x 1.25."),
    ("Own worker cost per hour", "Used for profit per client = (Basic+Allowances)/(Std days x hours per day)."),
    ("Sample data", "All names, companies, IDs and rates are fictional examples."),
]
r = 4
for a, b in lines:
    ws.cell(row=r, column=1, value=a).font = F_BOLD
    if b is None:
        ws.cell(row=r, column=1).font = Font(name=FONT, size=11, bold=True, color="1F3864")
    elif b:
        ws.cell(row=r, column=2, value=b).font = F_BASE
    r += 1
ws["A22"].fill = FILL_INPUT
ws["A22"].font = F_INPUT
ws["A25"].fill = FILL_RED
ws.column_dimensions["A"].width = 30
ws.column_dimensions["B"].width = 110

# ================================================================ Settings
ws = wb.create_sheet("Settings")
title(ws, "Settings", "Yellow cells are inputs. Change the payroll month here every month.")
settings = [
    ("Company Name", "Gulf Star Manpower Services LLC", None),
    ("Company TRN (VAT)", "100123456700003", None),
    ("Employer MOHRE ID (13 digits)", "0000001234567", None),
    ("Employer Bank Routing Code", "803320101", None),
    ("Payroll Month (1st day)", date(2026, 10, 1), DATE),
    ("Month End", "=EOMONTH(B8,0)", DATE),
    ("Salary Day Divisor", 30, None),
    ("Standard Working Days / Month", 26, None),
    ("Normal Hours / Day", 8, None),
    ("OT Multiplier", 1.25, "0.00"),
    ("VAT Rate", 0.05, "0.0%"),
    ("Document Expiry Alert (days)", 30, None),
]
for i, (k, v, fmt) in enumerate(settings):
    rr = 4 + i
    put(ws, rr, 1, k, bold=True)
    c = put(ws, rr, 2, v, fmt, inp=not str(v).startswith("="))
    if not str(v).startswith("="):
        c.fill = FILL_INPUT
ws.column_dimensions["A"].width = 32
ws.column_dimensions["B"].width = 34
for col, head, vals in (("D", "Trades", TRADES), ("F", "Source", ["Own", "Supplier"]),
                        ("H", "Pay Basis", ["Monthly", "Hourly"]),
                        ("J", "Deduction Types", DED_TYPES),
                        ("L", "Status", ["Active", "Inactive"])):
    c = ws[f"{col}4"]
    c.value = head
    c.font = F_HEAD
    c.fill = FILL_HEAD
    for i, v in enumerate(vals):
        put(ws, 5 + i, ws[f"{col}4"].column, v, inp=True)
    ws.column_dimensions[col].width = 16
ws["A17"] = "Note: OT multiplier 1.25 and divisor 30 follow common UAE practice - confirm with your HR / legal advisor."
ws["A17"].font = F_SUB

# ================================================================ Clients
clients = [
    ("C001", "Al Waha Contracting LLC", "100555111200003", "Dubai", 30),
    ("C002", "Blue Creek Facilities Mgmt", "100555222300003", "Abu Dhabi", 45),
    ("C003", "Sandline Construction Co.", "100555333400003", "Sharjah", 30),
    ("C004", "Marina Hospitality Group", "100555444500003", "Dubai", 60),
]
ws = wb.create_sheet("Clients")
title(ws, "Clients (Customers)", "One row per client. Client ID must be unique.")
headers(ws, [("Client ID", 11), ("Client Name", 32), ("TRN", 20), ("Emirate", 14),
             ("Payment Terms (days)", 14), ("Contact Person", 20), ("Phone", 16)])
for i, (cid, n, trn, em, terms) in enumerate(clients):
    rr = FIRST + i
    for j, v in enumerate([cid, n, trn, em, terms, "Mr. " + n.split()[0], "+971 4 000 00%02d" % i], 1):
        put(ws, rr, j, v, inp=True)

# ================================================================ Suppliers
suppliers = [
    ("S001", "Desert Falcon Manpower Supply", "100777111200003", 30),
    ("S002", "Emirates Link Labour Services", "100777222300003", 45),
    ("S003", "Oasis Workforce Solutions", "100777333400003", 30),
]
ws = wb.create_sheet("Suppliers")
title(ws, "Suppliers (Manpower Agencies)", "Companies that supply workers to you.")
headers(ws, [("Supplier ID", 11), ("Supplier Name", 32), ("TRN", 20),
             ("Payment Terms (days)", 14), ("Contact Person", 20), ("Bank IBAN", 28)])
for i, (sid, n, trn, terms) in enumerate(suppliers):
    rr = FIRST + i
    for j, v in enumerate([sid, n, trn, terms, "Mr. " + n.split()[0],
                           "AE07 0331 2345 6789 0123 45%d" % i], 1):
        put(ws, rr, j, v, inp=True)

# ================================================================ Rates
base_bill = {"Helper": 16, "Mason": 22, "Carpenter": 22, "Electrician": 26, "Plumber": 25,
             "Steel Fixer": 23, "Driver": 24, "Cleaner": 15, "Security Guard": 17}


def rate_sheet(name, key_label, ids, factor, sub):
    ws = wb.create_sheet(name)
    title(ws, name, sub)
    headers(ws, [("Key (auto)", 18), (key_label, 12), ("Trade", 16),
                 ("Normal Rate / hr (AED)", 14), ("OT Rate / hr (AED)", 14)])
    rr = FIRST
    for k, f in zip(ids, factor):
        for t in TRADES:
            put(ws, rr, 1, f'=IF(B{rr}="","",B{rr}&"|"&C{rr})')
            put(ws, rr, 2, k, inp=True)
            put(ws, rr, 3, t, inp=True)
            nr = round(base_bill[t] * f, 2)
            put(ws, rr, 4, nr, AED, inp=True)
            put(ws, rr, 5, round(nr * 1.25, 2), AED, inp=True)
            rr += 1
    for rr2 in range(rr, rr + 60):  # spare rows with key formula
        put(ws, rr2, 1, f'=IF(B{rr2}="","",B{rr2}&"|"&C{rr2})')
    add_list_validation(ws, f"C{FIRST}:C{rr + 60}", "=Settings!$D$5:$D$13")
    return ws


rate_sheet("ClientRates", "Client ID", [c[0] for c in clients], [1.0, 1.1, 0.95, 1.05],
           "Selling rate charged to each client, per trade.")
rate_sheet("SupplierRates", "Supplier ID", [s[0] for s in suppliers], [0.72, 0.75, 0.7],
           "Cost rate charged by each supplier, per trade.")

# ================================================================ Workers
first_names = ["Rahul", "Anwar", "Suresh", "Imran", "Ramesh", "Bilal", "Prakash", "Faisal",
               "Joseph", "Sanjay", "Asif", "Bikash", "Rohan", "Kamal", "Hari", "Jamal",
               "Mohan", "Tariq", "Ajith", "Shafi", "Ronel", "Marlon", "Dinesh", "Nabin"]
last_names = ["Kumar", "Ali", "Nair", "Khan", "Thapa", "Hossain", "Menon", "Ahmed",
              "Gurung", "Pillai", "Santos", "Reyes", "Shrestha", "Rahman", "Varghese"]
nationalities = ["India", "Pakistan", "Nepal", "Bangladesh", "Philippines"]

workers = []  # dicts
for i in range(40):  # own workers
    t = random.choice(TRADES)
    basis = "Hourly" if i % 8 == 7 else "Monthly"
    basic = random.choice([1200, 1400, 1500, 1800, 2000, 2500])
    workers.append(dict(id="W%04d" % (i + 1), src="Own", sup="", trade=t, basis=basis,
                        basic=basic if basis == "Monthly" else None,
                        allow=random.choice([300, 400, 500]) if basis == "Monthly" else None,
                        hr=random.choice([9, 10, 11]) if basis == "Hourly" else None))
for i in range(30):  # supplier workers
    workers.append(dict(id="SW%04d" % (i + 1), src="Supplier", sup=suppliers[i % 3][0],
                        trade=random.choice(TRADES), basis="", basic=None, allow=None, hr=None))

ws = wb.create_sheet("Workers")
title(ws, "Workers Master",
      "Own + supplier workers. Salary / bank fields are needed only for Own workers. Columns R-V are automatic.")
wcols = [("Worker ID", 10), ("Name", 22), ("Source", 10), ("Supplier ID", 10), ("Trade", 14),
         ("Nationality", 12), ("Passport No", 13), ("Labour Card / Person ID", 17),
         ("Visa Expiry", 13), ("Emirates ID Expiry", 13), ("Pay Basis", 10),
         ("Basic Salary / Month", 12), ("Allowances / Month", 12), ("Hourly Rate", 10),
         ("Bank Routing Code", 13), ("IBAN", 27), ("Status", 9),
         ("Days to Visa Expiry", 11), ("Days to EID Expiry", 11), ("Document Alert", 16),
         ("Own Seq (auto)", 9), ("Own Cost / hr", 10)]
headers(ws, wcols, calc_from=18)
W_MAX = 2000  # formula rows prepared for up to ~2000 workers
today = date(2026, 11, 1)
for i, w in enumerate(workers):
    rr = FIRST + i
    name = f"{random.choice(first_names)} {random.choice(last_names)}"
    exp_offset = random.choice([-10, 12, 25] + [random.randint(60, 700) for _ in range(12)])
    visa = date.fromordinal(today.toordinal() + exp_offset)
    eid = date.fromordinal(today.toordinal() + exp_offset + random.choice([0, 5, 40]))
    vals = [w["id"], name, w["src"], w["sup"] or None, w["trade"], random.choice(nationalities),
            "P%07d" % random.randint(1000000, 9999999),
            "%014d" % random.randint(10 ** 13, 10 ** 14 - 1) if w["src"] == "Own" else None,
            visa, eid, w["basis"] or None, w["basic"], w["allow"], w["hr"],
            random.choice(["803320101", "600310101", "402320106"]) if w["src"] == "Own" else None,
            ("AE%02d %04d %04d %04d %04d %03d" % (random.randint(10, 99), *[random.randint(1000, 9999) for _ in range(4)],
                                               random.randint(100, 999))) if w["src"] == "Own" else None,
            "Active"]
    for j, v in enumerate(vals, 1):
        put(ws, rr, j, v, DATE if j in (9, 10) else (AED if j in (12, 13, 14) else None), inp=True)
for rr in range(FIRST, FIRST + W_MAX):
    put(ws, rr, 18, f'=IF(I{rr}="","",I{rr}-TODAY())', NUM)
    put(ws, rr, 19, f'=IF(J{rr}="","",J{rr}-TODAY())', NUM)
    put(ws, rr, 20, f'=IF(A{rr}="","",IF(OR(AND(R{rr}<>"",R{rr}<0),AND(S{rr}<>"",S{rr}<0)),"EXPIRED",'
                    f'IF(OR(AND(R{rr}<>"",R{rr}<={S_ALERT}),AND(S{rr}<>"",S{rr}<={S_ALERT})),"EXPIRING SOON","OK")))')
    put(ws, rr, 21, f'=IF(AND(C{rr}="Own",Q{rr}="Active"),COUNTIFS($C${FIRST}:C{rr},"Own",$Q${FIRST}:Q{rr},"Active"),"")')
    put(ws, rr, 22, f'=IF(C{rr}<>"Own","",IF(K{rr}="Hourly",N{rr},(L{rr}+M{rr})/({S_STD}*{S_HRS})))', AED)
WL = FIRST + W_MAX - 1
add_list_validation(ws, f"C{FIRST}:C{WL}", "=Settings!$F$5:$F$6")
add_list_validation(ws, f"D{FIRST}:D{WL}", f"=Suppliers!$A${FIRST}:$A$200")
add_list_validation(ws, f"E{FIRST}:E{WL}", "=Settings!$D$5:$D$13")
add_list_validation(ws, f"K{FIRST}:K{WL}", "=Settings!$H$5:$H$6")
add_list_validation(ws, f"Q{FIRST}:Q{WL}", "=Settings!$L$5:$L$6")
ws.conditional_formatting.add(f"A{FIRST}:T{WL}", FormulaRule(formula=[f'$T{FIRST}="EXPIRED"'], fill=FILL_RED))
ws.conditional_formatting.add(f"A{FIRST}:T{WL}", FormulaRule(formula=[f'$T{FIRST}="EXPIRING SOON"'], fill=FILL_AMBER))
ws.auto_filter.ref = f"A{HEADER_ROW}:V{WL}"

# Worker lookup helper
WR = lambda col: f"Workers!${col}${FIRST}:${col}${WL}"  # noqa: E731

# ================================================================ Timesheet
ws = wb.create_sheet("Timesheet")
title(ws, "Timesheet (from Client)",
      "One row per worker per client per month. Enter columns A-H only; columns I-T calculate automatically.")
tcols = [("Month (1st day)", 12), ("Client ID", 9), ("Site / Project", 18), ("Worker ID", 10),
         ("Days Worked", 8), ("Normal Hours", 9), ("OT Hours", 8), ("Client Timesheet Ref", 14),
         ("Worker Name", 20), ("Source", 9), ("Supplier ID", 9), ("Trade", 13),
         ("Bill Rate / hr", 9), ("Bill OT Rate / hr", 9), ("Billing Amount (AED)", 13),
         ("Supplier Rate / hr", 9), ("Supplier OT Rate / hr", 9), ("Supplier Cost (AED)", 12),
         ("Own Worker Cost Est. (AED)", 12), ("Gross Profit (AED)", 12)]
headers(ws, tcols, calc_from=9)
sites = {"C001": ["Tower A - JVC", "Villa Project - Mirdif"], "C002": ["Yas Mall FM", "Khalifa City FM"],
         "C003": ["Warehouse - Sajja"], "C004": ["Hotel - Marina", "Hotel - Downtown"]}
rows = []
month = date(2026, 10, 1)
for w in workers:
    cid = random.choice(list(sites))
    days = random.randint(22, 26)
    if w["src"] == "Own" and random.random() < 0.15:  # split across 2 clients
        cid2 = random.choice([c for c in sites if c != cid])
        d1 = random.randint(8, 14)
        rows.append((month, cid, random.choice(sites[cid]), w["id"], d1, d1 * 8, random.choice([0, 6, 10]), f"TS-{cid}-2610"))
        rows.append((month, cid2, random.choice(sites[cid2]), w["id"], days - d1, (days - d1) * 8, random.choice([0, 4, 8]), f"TS-{cid2}-2610"))
    else:
        rows.append((month, cid, random.choice(sites[cid]), w["id"], days, days * 8,
                     random.choice([0, 0, 10, 16, 20, 26, 30]), f"TS-{cid}-2610"))
for w in workers[:12]:  # a few previous-month rows to show month filtering works
    rows.append((date(2026, 9, 1), "C001", "Tower A - JVC", w["id"], 25, 200, 12, "TS-C001-2609"))
T_MAX = 3000
for i, row in enumerate(rows):
    rr = FIRST + i
    for j, v in enumerate(row, 1):
        put(ws, rr, j, v, DATE if j == 1 else None, inp=True)
for rr in range(FIRST, FIRST + T_MAX):
    blank = f'$D{rr}=""'
    m = lambda col: f'INDEX({WR(col)},MATCH($D{rr},{WR("A")},0))'  # noqa: E731
    put(ws, rr, 9, f'=IF({blank},"",IFERROR({m("B")},"NOT FOUND"))')
    put(ws, rr, 10, f'=IF({blank},"",IFERROR({m("C")},""))')
    put(ws, rr, 11, f'=IF({blank},"",IFERROR(IF({m("D")}=0,"",{m("D")}),""))')
    put(ws, rr, 12, f'=IF({blank},"",IFERROR({m("E")},""))')
    ck = f'$B{rr}&"|"&$L{rr}'
    sk = f'$K{rr}&"|"&$L{rr}'
    put(ws, rr, 13, f'=IF({blank},"",IFERROR(INDEX(ClientRates!$D$5:$D$500,MATCH({ck},ClientRates!$A$5:$A$500,0)),0))', AED)
    put(ws, rr, 14, f'=IF({blank},"",IFERROR(INDEX(ClientRates!$E$5:$E$500,MATCH({ck},ClientRates!$A$5:$A$500,0)),0))', AED)
    put(ws, rr, 15, f'=IF({blank},"",F{rr}*M{rr}+G{rr}*N{rr})', AED)
    put(ws, rr, 16, f'=IF(OR({blank},$J{rr}<>"Supplier"),"",IFERROR(INDEX(SupplierRates!$D$5:$D$500,MATCH({sk},SupplierRates!$A$5:$A$500,0)),0))', AED)
    put(ws, rr, 17, f'=IF(OR({blank},$J{rr}<>"Supplier"),"",IFERROR(INDEX(SupplierRates!$E$5:$E$500,MATCH({sk},SupplierRates!$A$5:$A$500,0)),0))', AED)
    put(ws, rr, 18, f'=IF(OR({blank},$J{rr}<>"Supplier"),0,F{rr}*P{rr}+G{rr}*Q{rr})', AED)
    put(ws, rr, 19, f'=IF(OR({blank},$J{rr}<>"Own"),0,IFERROR((F{rr}+G{rr}*{S_OT})*{m("V")},0))', AED)
    put(ws, rr, 20, f'=IF({blank},"",O{rr}-R{rr}-S{rr})', AED)
TL = FIRST + T_MAX - 1
add_list_validation(ws, f"B{FIRST}:B{TL}", f"=Clients!$A${FIRST}:$A$300")
ws.conditional_formatting.add(f"I{FIRST}:I{TL}", FormulaRule(formula=[f'$I{FIRST}="NOT FOUND"'], fill=FILL_RED))
ws.conditional_formatting.add(f"M{FIRST}:M{TL}", FormulaRule(formula=[f'AND($D{FIRST}<>"",$M{FIRST}=0)'], fill=FILL_RED))
ws.conditional_formatting.add(f"P{FIRST}:P{TL}", FormulaRule(formula=[f'AND($J{FIRST}="Supplier",$P{FIRST}=0)'], fill=FILL_RED))
ws.auto_filter.ref = f"A{HEADER_ROW}:T{TL}"

TS = lambda col: f"Timesheet!${col}${FIRST}:${col}${TL}"  # noqa: E731
IN_MONTH = f'{TS("A")},">="&{S_START},{TS("A")},"<="&{S_END}'

# ================================================================ Deductions
ws = wb.create_sheet("Deductions")
title(ws, "Advances & Deductions (Own workers)", "Deduct in Month = the payroll month in which the amount is recovered.")
headers(ws, [("Date", 12), ("Worker ID", 10), ("Type", 14), ("Amount (AED)", 12),
             ("Deduct in Month (1st day)", 14), ("Notes", 30)])
deds = [(date(2026, 10, 3), "W0002", "Advance", 300, "Family emergency"),
        (date(2026, 10, 5), "W0005", "Advance", 500, ""),
        (date(2026, 10, 9), "W0011", "Fine", 100, "Late arrival x3"),
        (date(2026, 10, 12), "W0014", "Accommodation", 150, "Private room"),
        (date(2026, 10, 15), "W0020", "Advance", 400, ""),
        (date(2026, 10, 18), "W0027", "Other", 75, "Lost safety helmet"),
        (date(2026, 10, 21), "W0033", "Advance", 250, "")]
for i, (d, wid, t, a, n) in enumerate(deds):
    rr = FIRST + i
    for j, v in enumerate([d, wid, t, a, date(2026, 10, 1), n], 1):
        put(ws, rr, j, v, DATE if j in (1, 5) else (AED if j == 4 else None), inp=True)
add_list_validation(ws, f"C{FIRST}:C2000", "=Settings!$J$5:$J$8")
DD = lambda col: f"Deductions!${col}${FIRST}:${col}$2000"  # noqa: E731

# ================================================================ Payroll
ws = wb.create_sheet("Payroll")
title(ws, "Payroll - Own Workers", "Fully automatic. Lists every Active Own worker; values for the month in Settings!B8.")
pcols = [("Worker ID", 10), ("Name", 22), ("Trade", 13), ("Pay Basis", 9), ("Days Worked", 8),
         ("Normal Hours", 9), ("OT Hours", 8), ("Absent Days", 8), ("Basic Salary", 11),
         ("Allowances", 11), ("Fixed Earned", 11), ("Hourly Earnings", 11), ("OT Pay", 10),
         ("Gross Pay", 12), ("Advances", 10), ("Other Deductions", 11), ("Total Deductions", 11),
         ("NET PAY (AED)", 13), ("Bank Routing", 12), ("IBAN", 27), ("Labour Card / Person ID", 17)]
headers(ws, pcols, calc_from=1)
P_MAX = 1200
PL = FIRST + P_MAX - 1
put(ws, 3, 1, "TOTALS", bold=True).fill = FILL_TOTAL
put(ws, 3, 2, f'=COUNTIF(A{FIRST}:A{PL},"?*")&" workers"', bold=True).fill = FILL_TOTAL
for c in range(5, 19):
    col = get_column_letter(c)
    put(ws, 3, c, f"=SUM({col}{FIRST}:{col}{PL})", AED if c >= 9 else NUM, bold=True).fill = FILL_TOTAL
for rr in range(FIRST, FIRST + P_MAX):
    nb = f'$A{rr}=""'
    m = lambda col: f'INDEX({WR(col)},MATCH($A{rr},{WR("A")},0))'  # noqa: E731
    put(ws, rr, 1, f'=IFERROR(INDEX({WR("A")},MATCH(ROW()-{FIRST - 1},{WR("U")},0)),"")')
    put(ws, rr, 2, f'=IF({nb},"",{m("B")})')
    put(ws, rr, 3, f'=IF({nb},"",{m("E")})')
    put(ws, rr, 4, f'=IF({nb},"",{m("K")})')
    put(ws, rr, 5, f'=IF({nb},"",SUMIFS({TS("E")},{TS("D")},$A{rr},{IN_MONTH}))', NUM)
    put(ws, rr, 6, f'=IF({nb},"",SUMIFS({TS("F")},{TS("D")},$A{rr},{IN_MONTH}))', NUM)
    put(ws, rr, 7, f'=IF({nb},"",SUMIFS({TS("G")},{TS("D")},$A{rr},{IN_MONTH}))', NUM)
    put(ws, rr, 8, f'=IF({nb},"",IF(D{rr}="Monthly",MAX(0,{S_STD}-E{rr}),0))', NUM)
    put(ws, rr, 9, f'=IF({nb},"",N({m("L")}))', AED)
    put(ws, rr, 10, f'=IF({nb},"",N({m("M")}))', AED)
    put(ws, rr, 11, f'=IF({nb},"",IF(D{rr}="Monthly",ROUND(MAX(0,(I{rr}+J{rr})-(I{rr}+J{rr})/{S_DIV}*H{rr}),2),0))', AED)
    put(ws, rr, 12, f'=IF({nb},"",IF(D{rr}="Hourly",ROUND(F{rr}*N({m("N")}),2),0))', AED)
    put(ws, rr, 13, f'=IF({nb},"",ROUND(G{rr}*IF(D{rr}="Hourly",N({m("N")}),I{rr}/{S_DIV}/{S_HRS})*{S_OT},2))', AED)
    put(ws, rr, 14, f'=IF({nb},"",K{rr}+L{rr}+M{rr})', AED)
    put(ws, rr, 15, f'=IF({nb},"",SUMIFS({DD("D")},{DD("B")},$A{rr},{DD("C")},"Advance",{DD("E")},">="&{S_START},{DD("E")},"<="&{S_END}))', AED)
    put(ws, rr, 16, f'=IF({nb},"",SUMIFS({DD("D")},{DD("B")},$A{rr},{DD("E")},">="&{S_START},{DD("E")},"<="&{S_END})-O{rr})', AED)
    put(ws, rr, 17, f'=IF({nb},"",O{rr}+P{rr})', AED)
    put(ws, rr, 18, f'=IF({nb},"",N{rr}-Q{rr})', AED, bold=True)
    put(ws, rr, 19, f'=IF({nb},"",{m("O")}&"")')
    put(ws, rr, 20, f'=IF({nb},"",{m("P")}&"")')
    put(ws, rr, 21, f'=IF({nb},"",{m("H")}&"")')
ws.auto_filter.ref = f"A{HEADER_ROW}:U{PL}"
PR = lambda col: f"Payroll!${col}${FIRST}:${col}${PL}"  # noqa: E731

# ================================================================ Payslip
ws = wb.create_sheet("Payslip")
ws.column_dimensions["A"].width = 3
for col, wdt in zip("BCDE", (26, 30, 28, 18)):
    ws.column_dimensions[col].width = wdt
ws["B1"] = f"={S_COMPANY}"
ws["B1"].font = F_TITLE
ws["B2"] = '="PAYSLIP for "&TEXT(Settings!$B$8,"mmmm yyyy")'
ws["B2"].font = Font(name=FONT, size=12, bold=True)
put(ws, 4, 2, "Select Worker ID  >>", bold=True)
c = put(ws, 4, 3, "W0001", inp=True)
c.fill = FILL_INPUT
add_list_validation(ws, "C4", f"=Payroll!$A${FIRST}:$A${PL}")
pm = lambda col: f'=IFERROR(INDEX({PR(col)},MATCH($C$4,{PR("A")},0)),"")'  # noqa: E731
info = [("Employee Name", "B"), ("Trade", "C"), ("Pay Basis", "D"), ("Labour Card / Person ID", "U"),
        ("Bank IBAN", "T"), ("Days Worked", "E"), ("Absent Days", "H"), ("Normal Hours", "F"), ("OT Hours", "G")]
for i, (lab, col) in enumerate(info):
    put(ws, 6 + i, 2, lab, bold=True)
    put(ws, 6 + i, 3, pm(col))
ws["D6"] = "EARNINGS (AED)"
ws["D6"].font = F_HEAD
ws["D6"].fill = FILL_HEAD
ws["E6"].fill = FILL_HEAD
earn = [("Basic Salary", "I"), ("Allowances", "J"), ("Fixed Earned (after absence)", "K"),
        ("Hourly Earnings", "L"), ("Overtime Pay", "M"), ("GROSS PAY", "N")]
for i, (lab, col) in enumerate(earn):
    put(ws, 7 + i, 4, lab, bold=(lab == "GROSS PAY"))
    put(ws, 7 + i, 5, pm(col), AED, bold=(lab == "GROSS PAY"))
ws["D14"] = "DEDUCTIONS (AED)"
ws["D14"].font = F_HEAD
ws["D14"].fill = FILL_HEAD
ws["E14"].fill = FILL_HEAD
for i, (lab, col) in enumerate([("Salary Advance", "O"), ("Other Deductions", "P"), ("TOTAL DEDUCTIONS", "Q")]):
    put(ws, 15 + i, 4, lab, bold=(i == 2))
    put(ws, 15 + i, 5, pm(col), AED, bold=(i == 2))
put(ws, 19, 4, "NET PAY", bold=True).fill = FILL_TOTAL
c = put(ws, 19, 5, pm("R"), AED, bold=True)
c.fill = FILL_TOTAL
ws["B22"] = "Employee Signature: ____________________"
ws["D22"] = "Authorised Signatory: ____________________"
ws["B22"].font = ws["D22"].font = F_BASE
ws.print_area = "B1:E23"

# ================================================================ ClientInvoice
ws = wb.create_sheet("ClientInvoice")
ws.column_dimensions["A"].width = 3
for col, wdt in zip("BCDEFGHI", (18, 12, 12, 13, 12, 12, 13, 15)):
    ws.column_dimensions[col].width = wdt
ws["B1"] = f"={S_COMPANY}"
ws["B1"].font = F_TITLE
ws["B2"] = f'="TRN: "&{S_TRN}'
ws["B2"].font = F_BASE
ws["G1"] = "TAX INVOICE"
ws["G1"].font = Font(name=FONT, size=16, bold=True, color="1F3864")
put(ws, 4, 2, "Select Client ID >>", bold=True)
c = put(ws, 4, 3, "C001", inp=True)
c.fill = FILL_INPUT
add_list_validation(ws, "C4", f"=Clients!$A${FIRST}:$A$300")
CL = lambda col: f'IFERROR(INDEX(Clients!${col}$5:${col}$300,MATCH($C$4,Clients!$A$5:$A$300,0)),"")'  # noqa: E731
put(ws, 5, 2, "Bill To:", bold=True)
put(ws, 5, 3, "=" + CL("B"))
put(ws, 6, 2, "Client TRN:", bold=True)
put(ws, 6, 3, "=" + CL("C"))
put(ws, 4, 7, "Invoice No:", bold=True)
c = put(ws, 4, 8, '="INV-"&TEXT(Settings!$B$8,"yymm")&"-"&$C$4')
put(ws, 5, 7, "Invoice Date:", bold=True)
put(ws, 5, 8, f"={S_END}", DATE)
put(ws, 6, 7, "Due Date:", bold=True)
put(ws, 6, 8, f'=IFERROR({S_END}+{CL("E")},"")', DATE)
put(ws, 7, 7, "Period:", bold=True)
put(ws, 7, 8, '=TEXT(Settings!$B$8,"mmm yyyy")')
ih = ["Trade", "Workers", "Normal Hrs", "Rate / hr", "OT Hrs", "OT Rate", "Amount (AED)", ""]
for j, h in enumerate(ih[:7]):
    c = ws.cell(row=9, column=2 + j, value=h)
    c.font = F_HEAD
    c.fill = FILL_HEAD
    c.alignment = Alignment(horizontal="center")
for i, t in enumerate(TRADES):
    rr = 10 + i
    crit = f'{TS("B")},$C$4,{TS("L")},$B{rr},{IN_MONTH}'
    put(ws, rr, 2, f"=Settings!$D${5 + i}")
    put(ws, rr, 3, f"=COUNTIFS({crit})", NUM)
    put(ws, rr, 4, f"=SUMIFS({TS('F')},{crit})", NUM)
    put(ws, rr, 5, f'=IFERROR(INDEX(ClientRates!$D$5:$D$500,MATCH($C$4&"|"&$B{rr},ClientRates!$A$5:$A$500,0)),0)', AED)
    put(ws, rr, 6, f"=SUMIFS({TS('G')},{crit})", NUM)
    put(ws, rr, 7, f'=IFERROR(INDEX(ClientRates!$E$5:$E$500,MATCH($C$4&"|"&$B{rr},ClientRates!$A$5:$A$500,0)),0)', AED)
    put(ws, rr, 8, f"=SUMIFS({TS('O')},{crit})", AED)
last = 10 + len(TRADES) - 1
put(ws, last + 2, 7, "Sub Total", bold=True)
put(ws, last + 2, 8, f"=SUM(H10:H{last})", AED, bold=True)
put(ws, last + 3, 7, '="VAT "&TEXT(Settings!$B$14,"0%")', bold=True)
put(ws, last + 3, 8, f"=ROUND(H{last + 2}*{S_VAT},2)", AED)
put(ws, last + 4, 7, "TOTAL (AED)", bold=True).fill = FILL_TOTAL
c = put(ws, last + 4, 8, f"=H{last + 2}+H{last + 3}", AED, bold=True)
c.fill = FILL_TOTAL
put(ws, last + 2, 2, "Total hours:", bold=True)
put(ws, last + 2, 3, f"=SUM(D10:D{last})+SUM(F10:F{last})", NUM)
ws.cell(row=last + 6, column=2, value="Detailed worker-wise timesheet: filter the Timesheet sheet by this Client ID and attach.").font = F_SUB
ws.print_area = f"B1:H{last + 7}"

# ================================================================ SupplierStatement
ws = wb.create_sheet("SupplierStatement")
title(ws, "Supplier Payment Statement", "Amount payable to each supplier for the month in Settings!B8.")
headers(ws, [("Supplier ID", 11), ("Supplier Name", 30), ("Worker Lines", 9), ("Days", 9),
             ("Normal Hours", 10), ("OT Hours", 9), ("Amount (AED)", 13), ("VAT (AED)", 11),
             ("Total Payable (AED)", 14), ("Payment Terms", 9), ("Due Date", 12)], calc_from=1)
for i in range(50):
    rr = FIRST + i
    nb = f'$A{rr}=""'
    crit = f'{TS("K")},$A{rr},{IN_MONTH}'
    put(ws, rr, 1, f'=IF(Suppliers!A{rr}="","",Suppliers!A{rr})')
    put(ws, rr, 2, f'=IF({nb},"",Suppliers!B{rr})')
    put(ws, rr, 3, f'=IF({nb},"",COUNTIFS({crit}))', NUM)
    put(ws, rr, 4, f'=IF({nb},"",SUMIFS({TS("E")},{crit}))', NUM)
    put(ws, rr, 5, f'=IF({nb},"",SUMIFS({TS("F")},{crit}))', NUM)
    put(ws, rr, 6, f'=IF({nb},"",SUMIFS({TS("G")},{crit}))', NUM)
    put(ws, rr, 7, f'=IF({nb},"",SUMIFS({TS("R")},{crit}))', AED)
    put(ws, rr, 8, f'=IF({nb},"",ROUND(G{rr}*{S_VAT},2))', AED)
    put(ws, rr, 9, f'=IF({nb},"",G{rr}+H{rr})', AED, bold=True)
    put(ws, rr, 10, f'=IF({nb},"",Suppliers!D{rr})')
    put(ws, rr, 11, f'=IF({nb},"",{S_END}+J{rr})', DATE)
put(ws, 3, 1, "TOTALS", bold=True).fill = FILL_TOTAL
for c in range(3, 10):
    col = get_column_letter(c)
    put(ws, 3, c, f"=SUM({col}{FIRST}:{col}{FIRST + 49})", AED if c >= 7 else NUM, bold=True).fill = FILL_TOTAL

# ================================================================ WPS_SIF
ws = wb.create_sheet("WPS_SIF")
title(ws, "WPS Salary Information File (SIF) - layout sample",
      "EDR rows per employee + SCR control row. Save as CSV named per bank rules. Verify field formats with your bank/exchange house.")
put(ws, 3, 1, "SCR", bold=True)
scr = [f"={S_MOHRE}", f"={S_BANK}", '=TEXT(TODAY(),"yyyy-mm-dd")', '=TEXT(NOW(),"hhmm")',
       '=TEXT(Settings!$B$8,"mmyyyy")', f'=COUNTIF(A{FIRST}:A{FIRST + 1199},"EDR")',
       f"=SUM(H{FIRST}:H{FIRST + 1199})+SUM(I{FIRST}:I{FIRST + 1199})", "AED", f"={S_COMPANY}"]
for j, f in enumerate(scr, 2):
    put(ws, 3, j, f, AED if j == 8 else None, bold=True).fill = FILL_TOTAL
headers(ws, [("Record", 7), ("Person ID / Labour Card", 18), ("Agent Routing Code", 13), ("IBAN", 27),
             ("Pay Start", 12), ("Pay End", 12), ("Days", 7), ("Fixed Salary", 12), ("Variable Salary", 12),
             ("Leave Days", 9)], calc_from=1)
for i in range(1200):
    rr = FIRST + i
    pr = FIRST + i
    ok = f'AND(Payroll!$A{pr}<>"",N(Payroll!$R{pr})>0)'
    put(ws, rr, 1, f'=IF({ok},"EDR","")')
    put(ws, rr, 2, f'=IF($A{rr}="","",Payroll!U{pr})')
    put(ws, rr, 3, f'=IF($A{rr}="","",Payroll!S{pr})')
    put(ws, rr, 4, f'=IF($A{rr}="","",SUBSTITUTE(Payroll!T{pr}," ",""))')
    put(ws, rr, 5, f'=IF($A{rr}="","",TEXT({S_START},"yyyy-mm-dd"))')
    put(ws, rr, 6, f'=IF($A{rr}="","",TEXT({S_END},"yyyy-mm-dd"))')
    put(ws, rr, 7, f'=IF($A{rr}="","",DAY({S_END}))')
    put(ws, rr, 8, f'=IF($A{rr}="","",ROUND(Payroll!R{pr}-Payroll!L{pr}-Payroll!M{pr},2))', AED)
    put(ws, rr, 9, f'=IF($A{rr}="","",ROUND(Payroll!L{pr}+Payroll!M{pr},2))', AED)
    put(ws, rr, 10, f'=IF($A{rr}="","",0)')
ws.freeze_panes = "A5"

# ================================================================ Dashboard
ws = wb.create_sheet("Dashboard", 1)
title(ws, "Dashboard", "")
ws["A2"] = '="Month: "&TEXT(Settings!$B$8,"mmmm yyyy")'
ws["A3"] = "(change month in Settings!B8)"
ws["A3"].font = F_SUB
ws["A2"].font = Font(name=FONT, size=11, bold=True, color="2E75B6")
ws.column_dimensions["A"].width = 30
for col, wdt in zip("BCDEFGHIJ", (16, 4, 14, 26, 12, 14, 14, 14, 14)):
    ws.column_dimensions[col].width = wdt
kpis = [
    ("Client Billing (excl. VAT)", f"=SUMIFS({TS('O')},{IN_MONTH})", AED),
    ("Supplier Cost (excl. VAT)", f"=SUMIFS({TS('R')},{IN_MONTH})", AED),
    ("Own Workers Payroll (Gross)", "=Payroll!$N$3", AED),
    ("GROSS PROFIT", "=B4-B5-B6", AED),
    ("Gross Margin %", "=IF(B4=0,0,B7/B4)", "0.0%"),
    ("VAT Output (on billing)", f"=ROUND(B4*{S_VAT},2)", AED),
    ("", None, None),
    ("Total Billable Hours", f"=SUMIFS({TS('F')},{IN_MONTH})+SUMIFS({TS('G')},{IN_MONTH})", NUM),
    ("Timesheet Lines (this month)", f'=COUNTIFS({IN_MONTH})', NUM),
    ("Active Own Workers", f'=COUNTIFS({WR("C")},"Own",{WR("Q")},"Active")', NUM),
    ("Active Supplier Workers", f'=COUNTIFS({WR("C")},"Supplier",{WR("Q")},"Active")', NUM),
    ("Net Salary to Pay (WPS)", "=Payroll!$R$3", AED),
    ("Payable to Suppliers (incl. VAT)", "=SupplierStatement!$I$3", AED),
    ("", None, None),
    ("Documents EXPIRED", f'=COUNTIF({WR("T")},"EXPIRED")', NUM),
    ("Documents EXPIRING SOON", f'=COUNTIF({WR("T")},"EXPIRING SOON")', NUM),
]
for i, (lab, f, fmt) in enumerate(kpis):
    rr = 4 + i
    if not lab:
        continue
    put(ws, rr, 1, lab, bold=True)
    c = put(ws, rr, 2, f, fmt, bold=True)
    if lab in ("GROSS PROFIT", "Gross Margin %"):
        c.fill = FILL_TOTAL
ws.conditional_formatting.add("B18:B19", FormulaRule(formula=["B18>0"], fill=FILL_RED))
heads = ["Client ID", "Client Name", "Lines", "Billing (AED)", "Supplier Cost", "Own Cost Est.", "Profit (AED)", "Margin %"]
for j, h in enumerate(heads):
    c = ws.cell(row=3, column=4 + j, value=h)
    c.font = F_HEAD
    c.fill = FILL_HEAD
    c.alignment = Alignment(horizontal="center", wrap_text=True)
ws.cell(row=2, column=4, value="Profit by Client (this month)").font = F_BOLD
for i in range(30):
    rr = 4 + i
    cr = FIRST + i
    nb = f'$D{rr}=""'
    crit = f'{TS("B")},$D{rr},{IN_MONTH}'
    put(ws, rr, 4, f'=IF(Clients!A{cr}="","",Clients!A{cr})')
    put(ws, rr, 5, f'=IF({nb},"",Clients!B{cr})')
    put(ws, rr, 6, f'=IF({nb},"",COUNTIFS({crit}))', NUM)
    put(ws, rr, 7, f'=IF({nb},"",SUMIFS({TS("O")},{crit}))', AED)
    put(ws, rr, 8, f'=IF({nb},"",SUMIFS({TS("R")},{crit}))', AED)
    put(ws, rr, 9, f'=IF({nb},"",SUMIFS({TS("S")},{crit}))', AED)
    put(ws, rr, 10, f'=IF({nb},"",G{rr}-H{rr}-I{rr})', AED, bold=True)
    put(ws, rr, 11, f'=IF(OR({nb},N(G{rr})=0),"",J{rr}/G{rr})', "0.0%")
ws.column_dimensions["K"].width = 10
ws["A22"] = "Note: 'Own Cost Est.' per client uses each own worker's cost/hour (Workers col V); Gross Profit at left uses actual payroll."
ws["A22"].font = F_SUB

# ================================================================ HOME (menu UI)
ws = wb.create_sheet("HOME", 0)
ws.sheet_view.showGridLines = False
ws.sheet_view.zoomScale = 110
for col in "ABCDEFGHIJKLM":
    ws.column_dimensions[col].width = 3 if col in "AEIM" else 15
ws.column_dimensions["A"].width = 2
ws["B1"] = f"={S_COMPANY}"
ws["B1"].font = Font(name=FONT, size=20, bold=True, color="1F3864")
ws["B2"] = '="Manpower Management System   |   Payroll month: "&TEXT(Settings!$B$8,"mmmm yyyy")'
ws["B2"].font = Font(name=FONT, size=11, color="595959")

# KPI strip
kpi_tiles = [("Client Billing (AED)", "=Dashboard!$B$4", AED), ("Gross Profit (AED)", "=Dashboard!$B$7", AED),
             ("Net Salary to Pay - WPS (AED)", "=Dashboard!$B$15", AED)]
for i, (lab, f, fmt) in enumerate(kpi_tiles):
    col = [2, 6, 10][i]
    for rr in (4, 5):
        ws.merge_cells(start_row=rr, start_column=col, end_row=rr, end_column=col + 2)
        for cc in range(col, col + 3):
            ws.cell(row=rr, column=cc).fill = PatternFill("solid", fgColor="F2F2F2")
    a = ws.cell(row=4, column=col, value=lab)
    a.font = Font(name=FONT, size=9, color="595959")
    a.alignment = Alignment(horizontal="center")
    v = ws.cell(row=5, column=col, value=f)
    v.font = Font(name=FONT, size=16, bold=True, color="1F3864")
    v.number_format = fmt
    v.alignment = Alignment(horizontal="center")
alert = ws.cell(row=6, column=2, value='=IF(Dashboard!$B$18+Dashboard!$B$19=0,"All visas / Emirates IDs OK",'
                '"WARNING: "&Dashboard!$B$18&" documents EXPIRED, "&Dashboard!$B$19&" expiring soon - open Workers")')
alert.font = Font(name=FONT, size=10, bold=True, color="C00000")
alert.hyperlink = "#'Workers'!A1"
ws.row_dimensions[5].height = 28

# Menu tiles: (section, colour, [(label, sheet, hint)])
menu = [
    ("1. ENTER DATA", "2E75B6", [("Timesheet", "Timesheet", "Client hours each month"),
                                 ("Deductions", "Deductions", "Advances, fines"),
                                 ("Workers", "Workers", "Add / edit workers")]),
    ("2. MASTER DATA", "7030A0", [("Clients", "Clients", "Customers"),
                                  ("Suppliers", "Suppliers", "Manpower agencies"),
                                  ("Client Rates", "ClientRates", "Selling rate / hr")]),
    ("", "7030A0", [("Supplier Rates", "SupplierRates", "Cost rate / hr"),
                    ("Settings", "Settings", "Month, OT, VAT"),
                    ("Help", "README", "How it works")]),
    ("3. REPORTS & PRINT", "548235", [("Dashboard", "Dashboard", "Profit & KPIs"),
                                      ("Payroll", "Payroll", "Salary sheet"),
                                      ("Payslip", "Payslip", "Print payslip")]),
    ("", "548235", [("Client Invoice", "ClientInvoice", "Tax invoice"),
                    ("Supplier Payment", "SupplierStatement", "Amount payable"),
                    ("WPS File", "WPS_SIF", "Bank salary file")]),
]
row = 8
for section, colour, items in menu:
    if section:
        ws.cell(row=row, column=2, value=section).font = Font(name=FONT, size=11, bold=True, color=colour)
        row += 1
    for i, (label, sheet, hint) in enumerate(items):
        col = [2, 6, 10][i]
        ws.merge_cells(start_row=row, start_column=col, end_row=row + 1, end_column=col + 2)
        c = ws.cell(row=row, column=col, value=label)
        c.hyperlink = f"#'{sheet}'!A1"
        c.font = Font(name=FONT, size=13, bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center")
        for rr in (row, row + 1):
            for cc in range(col, col + 3):
                ws.cell(row=rr, column=cc).fill = PatternFill("solid", fgColor=colour)
        h = ws.cell(row=row + 2, column=col, value=hint)
        h.font = Font(name=FONT, size=8, italic=True, color="7F7F7F")
    row += 4
ws.cell(row=row, column=2, value="Click any tile to open. Every sheet has a  « HOME  link at the top right to come back.").font = F_SUB

# "« HOME" link on every other sheet (placed right of the title so it doesn't hide it)
home_pos = {"Payslip": "E1", "ClientInvoice": "I1", "Dashboard": "D1", "Settings": "D1"}


def home_cell(s):
    if s.title in home_pos:
        return home_pos[s.title]
    need = len(str(s["A1"].value or "")) * 1.6 + 4
    acc, col = 0, 1
    while acc < need:
        acc += s.column_dimensions[get_column_letter(col)].width or 8.43
        col += 1
    return f"{get_column_letter(col)}1"


for s in wb.worksheets:
    if s.title == "HOME":
        continue
    c = s[home_cell(s)]
    c.value = "« HOME"
    c.hyperlink = "#'HOME'!A1"
    c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF", underline="single")
    c.fill = PatternFill("solid", fgColor="C00000")
    c.alignment = Alignment(horizontal="center", vertical="center")
wb.active = 0

# ---------------------------------------------------------------- finish
for s in wb.worksheets:
    s.sheet_view.showGridLines = s.title in ("Settings",)
    s.sheet_properties.tabColor = {"HOME": "C00000", "README": "7F7F7F", "Dashboard": "C00000", "Payroll": "548235",
                                   "Payslip": "548235", "ClientInvoice": "548235",
                                   "SupplierStatement": "548235", "WPS_SIF": "548235"}.get(s.title, "1F3864")
wb.save(OUT)
print("saved", OUT)
